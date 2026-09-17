import json,re,math
from docx import Document
from docx.shared import Pt
S='/tmp/claude-0/-home-user-imageforge/762fa79f-125c-58b7-bf46-ebb266cfbbf8/scratchpad/'
OUT='/home/user/imageforge/docs/robert-pattinson/'
L=open(S+'packets.txt').read().split('\n')
d=json.load(open(S+'v2.json')); clips=d['clips']; changes=d['changes']; total=d['total']
HEADERS=['ATTACH THESE REFERENCES','TONE/PERFORMANCE','BACKGROUND CONTEXT ONLY / DO NOT DEPICT','OPENING STATE','CURRENT CLIP — SHOOT ONLY THIS','EXACT SPOKEN WORDS','ENDING STATE','EXCLUSIONS']
SAYTRAIL='Speaker labels are not spoken. Say the words exactly, including incomplete clauses; do not vocalize punctuation or draft annotation (?). Only a label marked VO indicates narration.'
def saylines(s):
    parts=re.split(r'(?=(?:Woman VO|Robert VO|Woman|Robert|Pedestrian|Dispatcher|Guard|PA voice|Waitress): )',s)
    return [p.strip() for p in parts if p.strip()]
def prompt_text(c):
    o=[f"CLIP {c['id']} — {c['sec']} SECONDS MAXIMUM", c['style'].strip()]
    for h in HEADERS:
        if h not in c: continue
        body=c[h].strip()
        if h=='EXACT SPOKEN WORDS':
            ls=saylines(c['say'])
            body='\n'.join(ls)+'\n'+SAYTRAIL if ls else 'None. No improvised dialogue or voiceover.'
        # reinsert line breaks the docx flattened: refs "R01: …R02: …", and sentences after "Start already"/"Allow a short"/"Hold this"
        body=re.sub(r'(?<=\.)(?=R\d\d: )','\n',body)
        body=re.sub(r'(?<=\.)(?=Use the actual attached|Start already|Allow a short|Hold this state|This context changes|Speaker labels)','\n',body)
        o.append(h); o.append(body)
    return '\n'.join(o)
# ---- header (decisions) patched
HEAD=L[:85]
head=[]
for ln in HEAD:
    if ln.startswith('Package:'):
        ln=f"Package: {len(clips)} clips, 57 reference records, approximately {total//60} minutes {total%60} seconds of planned footage before trimming. No clip exceeds 15 seconds. Every clip has a reference-matrix row, a complete prompt and a production-tracker row. The long confession is preserved word for word; it is split at sentence boundaries only (v2), never mid-clause."
    ln=ln.replace('Base speech packets use no more than 26 words per 15 seconds, with a few short exchanges up to 33 including speaker labels.','Base speech packets carry up to about 40 words per 15 seconds (roughly 160 words a minute, the pace the source directs for her: “speaking more quickly”, “in a rush”), and a clip is cut only at the end of a sentence.')
    ln=ln.replace('Robert Pattinson Seedance clip packets and reference tracker','Robert Pattinson Seedance clip packets and reference tracker — v2 (2026-09-17: fidelity fixes, see “What v2 changed”)')
    ln=ln.replace('selected opening car prose is retained as voiceover','selected opening car prose is retained as the woman’s voiceover, and seven lines of Robert’s chapter-A prose are proposed as his (marked VO PROPOSED)')
    if ln.strip()=='How to use each packet':
        head+=['What v2 changed (2026-09-17)','Each row is a fidelity correction against the two source chapters. The clips named are v2 ids.']
        for cid,what in changes: head+=[cid,what]
        head+=['Ambiguities still open for the author','Which lines carry the mocking eyelash wish (v2: Robert) · whether “Yes” on the ambulance call is Robert or the dispatcher (v2: dispatcher) · who says “Hey - you!” (v2: Robert, softening at once) · “It was like I was bullying her, but I wasn’t” sits outside the quotation marks (v2: hers) · “The girl rolled her eyes” — the woman as a child or Robert watching her now (v2: her now) · raspberries then apples · “something in mean” (v2: “something mean”) · Paige in A, Stella in B · what happens to the injured man · how they reach the Reel Inn · where the cap and sweatpants go before the long black coat · the first candle’s extinction before she lights it again.']
    head.append(ln)
# new decision rows
ledger=L[85:370]
ledger=[x.replace('Red heart, white doily, Robert magazine cutouts;','White doily, Robert magazine cutouts (the source’s “awful doilie thing” — no red heart is specified);').replace('same heart height','same valentine height') for x in ledger]
# ---- clip index
index=['Clip index and reference matrix','Every row is one base-edit clip. Ref IDs expand in the ledger; they are requirements, not a claim the service can attach this many separate images. Consolidate into an approved opening frame if needed.','Clip','Source','Seconds','Beat','References','Gate']
for c in clips: index+=[c['id'],c['src'],str(c['sec']),c['title'],c['refs'],c['gate']]
# ---- packets
packets=['Generation packets','The sequence below is the edit order. V recollection clips are placed after the first watcher clip and before return to present-day cafe. All prompts state present opening/ending conditions; no prior clip is assumed to be remembered by the model.']
for c in clips: packets+=[f"{c['id']} {c['title']}",f"Source: {c['src']}. Target: {c['sec']} seconds. Gate: {c['gate']}.",prompt_text(c)]
tracker=['Production tracker rows','Fill one row per approved take; retain rejected take filenames separately. QA covers exact dialogue, identity, left/right geography, props, weather, duration and ending state. No outputs exist yet.','Clip','Source gate','References','Output and take','QA','Closing frame','Decision']
for c in clips: tracker+=[c['id'],c['gate'],'Not visually approved','—','Not run','—','Not generated']
APP=L[3362:]
ALL=head+ledger+index+packets+tracker+APP
# ---- md
md=[]
for ln in ALL:
    if ln.startswith('CLIP ') and '\n' in ln: md.append('```\n'+ln+'\n```')
    elif re.fullmatch(r'[A-Z]\d{3} .+',ln) and not ln.startswith('CLIP'): md.append('### '+ln)
    elif ln in ('Reference image ledger','Clip index and reference matrix','Generation packets','Production tracker rows','Original source text','What v2 changed (2026-09-17)','Before generating','Source order and decisions','How to use each packet','Ambiguities still open for the author'): md.append('## '+ln)
    else: md.append(ln)
open(OUT+'seedance-packets-v2.md','w').write('# Robert Pattinson Seedance Mini Production Package — v2\n\n'+'\n\n'.join(x for x in md if x.strip()))
# ---- docx
doc=Document(); st=doc.styles['Normal']; st.font.size=Pt(10)
doc.add_heading('Robert Pattinson Seedance Mini Production Package — v2',0)
for ln in ALL:
    if not ln.strip(): continue
    if ln.startswith('CLIP ') and '\n' in ln:
        p=doc.add_paragraph(); r=p.add_run(ln); r.font.name='Courier New'; r.font.size=Pt(8.5)
    elif re.fullmatch(r'[A-Z]\d{3} .+',ln): doc.add_heading(ln,3)
    elif ln in ('Reference image ledger','Clip index and reference matrix','Generation packets','Production tracker rows','Original source text','What v2 changed (2026-09-17)','Before generating','Source order and decisions','How to use each packet','Ambiguities still open for the author'): doc.add_heading(ln,1)
    else: doc.add_paragraph(ln)
doc.save(OUT+'Robert_Pattinson_Seedance_Packets_v2.docx')
print(len(ALL),'lines; clips',len(clips))
